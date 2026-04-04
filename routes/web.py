from flask import Blueprint, render_template, request, jsonify, flash, redirect, url_for, session, send_file
import os
import io
from werkzeug.utils import secure_filename
from services.gpt_service import get_gpt_response, analyze_image, generate_image, should_generate_image, is_image_followup, extract_image_prompt_from_history, get_voice_response
from services.pdf_service import extract_text_from_pdf
from services.whisper_service import transcribe_audio
from services.tts_service import text_to_speech, tts_to_bytes
from services.file_generator import should_generate_file, generate_pdf, generate_word, generate_pptx
from database import save_conversation, get_conversations, get_conversation_history, create_chat_session, get_chat_sessions, get_session_messages, update_session_title, get_user_by_id, delete_chat_session

def login_required(f):
    """Décorateur pour protéger les routes"""
    from functools import wraps
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect(url_for('auth.login'))
        return f(*args, **kwargs)
    return decorated_function

web_bp = Blueprint('web', __name__)

UPLOAD_FOLDER = 'uploads'
ALLOWED_EXTENSIONS = {
    # Images
    'png', 'jpg', 'jpeg', 'gif', 'webp', 'bmp',
    # Audio
    'mp3', 'wav', 'ogg', 'm4a', 'flac', 'webm',
    # Documents Adobe / texte
    'pdf', 'txt',
    # Microsoft Office
    'doc', 'docx', 'xls', 'xlsx', 'ppt', 'pptx',
    # Données
    'csv', 'json', 'xml',
    # OpenDocument
    'odt', 'ods', 'odp',
}

# Créer le dossier uploads s'il n'existe pas
os.makedirs(UPLOAD_FOLDER, exist_ok=True)

def allowed_file(filename):
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

@web_bp.route('/')
@web_bp.route('/chat/<session_id>')
@login_required
def index(session_id=None):
    """Page d'accueil avec interface de chat"""
    user_id = session['user_id']
    user = get_user_by_id(user_id)
    
    if not session_id:
        session_id = create_chat_session(user_id)
        return redirect(f'/chat/{session_id}')
    
    sessions = get_chat_sessions(user_id)
    messages = get_session_messages(session_id)
    return render_template('index.html', sessions=sessions, current_session=session_id, messages=messages, user=user)

@web_bp.route('/chat/<session_id>', methods=['POST'])
@login_required
def chat(session_id):
    """Traite les messages et fichiers depuis l'interface web"""
    try:
        message = request.form.get('message', '').strip()
        file = request.files.get('file')
        
        response_text = ""
        message_type = "texte"
        question = message
        sources = []  # liste de {title, url} à afficher sous la réponse
        
        # Récupérer l'historique de conversation
        conversation_history = get_conversation_history(session_id, 6)
        
        if file and file.filename and allowed_file(file.filename):
            # Sauvegarder le fichier
            filename = secure_filename(file.filename)
            filepath = os.path.join(UPLOAD_FOLDER, filename)
            file.save(filepath)
            
            try:
                # Traitement selon le type de fichier
                if filename.lower().endswith(('.png', '.jpg', '.jpeg', '.gif')):
                    message_type = "image"
                    import base64
                    with open(filepath, 'rb') as img_file:
                        image_base64 = base64.b64encode(img_file.read()).decode('utf-8')
                    from services.gpt_service import analyze_image_base64
                    prompt = message if message else "Décris cette image en détail."
                    response_text = analyze_image_base64(image_base64, prompt)
                    question = f"Image: {filename}" + (f" + Question: {message}" if message else "")
                    
                elif filename.lower().endswith('.pdf'):
                    message_type = "pdf"
                    pdf_text = extract_text_from_pdf(filepath)
                    if pdf_text and not pdf_text.startswith("Erreur"):
                        prompt = f"Analyse ce document PDF:\n\n{pdf_text}"
                        if message:
                            prompt += f"\n\nQuestion spécifique: {message}"
                        response_text, _ = get_gpt_response(prompt, max_tokens=1500, use_search=False, conversation_history=conversation_history)
                        question = f"PDF: {filename}" + (f" + Question: {message}" if message else "")
                    else:
                        response_text = pdf_text

                elif filename.lower().endswith(('.mp3', '.wav', '.ogg', '.m4a', '.flac', '.webm')):
                    message_type = "audio"
                    transcription = transcribe_audio(filepath)
                    if transcription and not transcription.startswith("Erreur"):
                        prompt = f"Réponds à cette transcription audio: {transcription}"
                        if message:
                            prompt += f"\n\nContext/Question: {message}"
                        response_text, _ = get_gpt_response(prompt, max_tokens=1500, use_search=False, conversation_history=conversation_history)
                        question = f"Audio: {filename}" + (f" + Question: {message}" if message else "")
                    else:
                        response_text = transcription

                elif filename.lower().endswith(('.docx', '.doc')):
                    message_type = "word"
                    try:
                        from docx import Document as _DocX
                        doc = _DocX(filepath)
                        doc_text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
                        prompt = f"Analyse ce document Word:\n\n{doc_text[:8000]}"
                        if message:
                            prompt += f"\n\nQuestion spécifique: {message}"
                        response_text, _ = get_gpt_response(prompt, max_tokens=1500, use_search=False, conversation_history=conversation_history)
                        question = f"Word: {filename}" + (f" + Question: {message}" if message else "")
                    except Exception as e:
                        response_text = f"Impossible de lire le fichier Word: {str(e)}"

                elif filename.lower().endswith(('.xlsx', '.xls')):
                    message_type = "excel"
                    try:
                        import openpyxl as _xl
                        wb = _xl.load_workbook(filepath, read_only=True, data_only=True)
                        rows = []
                        for sheet in wb.worksheets[:3]:
                            rows.append(f"[Feuille: {sheet.title}]")
                            for row in sheet.iter_rows(max_row=100, values_only=True):
                                if any(c is not None for c in row):
                                    rows.append('\t'.join(str(c) if c is not None else '' for c in row))
                        sheet_text = '\n'.join(rows)[:8000]
                        prompt = f"Analyse ce fichier Excel:\n\n{sheet_text}"
                        if message:
                            prompt += f"\n\nQuestion spécifique: {message}"
                        response_text, _ = get_gpt_response(prompt, max_tokens=1500, use_search=False, conversation_history=conversation_history)
                        question = f"Excel: {filename}" + (f" + Question: {message}" if message else "")
                    except Exception as e:
                        response_text = f"Impossible de lire le fichier Excel: {str(e)}"

                elif filename.lower().endswith(('.pptx', '.ppt')):
                    message_type = "powerpoint"
                    try:
                        from pptx import Presentation as _Prs
                        prs = _Prs(filepath)
                        slides_text = []
                        for i, slide in enumerate(prs.slides, 1):
                            texts = [sh.text for sh in slide.shapes if hasattr(sh, 'text') and sh.text.strip()]
                            if texts:
                                slides_text.append(f"Slide {i}: " + ' | '.join(texts))
                        pptx_text = '\n'.join(slides_text)[:8000]
                        prompt = f"Analyse cette présentation PowerPoint:\n\n{pptx_text}"
                        if message:
                            prompt += f"\n\nQuestion spécifique: {message}"
                        response_text, _ = get_gpt_response(prompt, max_tokens=1500, use_search=False, conversation_history=conversation_history)
                        question = f"PowerPoint: {filename}" + (f" + Question: {message}" if message else "")
                    except Exception as e:
                        response_text = f"Impossible de lire le fichier PowerPoint: {str(e)}"

                elif filename.lower().endswith('.csv'):
                    message_type = "csv"
                    try:
                        import csv as _csv
                        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                            reader = _csv.reader(f)
                            rows = ['\t'.join(row) for _, row in zip(range(100), reader)]
                        csv_text = '\n'.join(rows)[:8000]
                        prompt = f"Analyse ce fichier CSV:\n\n{csv_text}"
                        if message:
                            prompt += f"\n\nQuestion spécifique: {message}"
                        response_text, _ = get_gpt_response(prompt, max_tokens=1500, use_search=False, conversation_history=conversation_history)
                        question = f"CSV: {filename}" + (f" + Question: {message}" if message else "")
                    except Exception as e:
                        response_text = f"Impossible de lire le fichier CSV: {str(e)}"

                elif filename.lower().endswith(('.txt', '.json', '.xml')):
                    message_type = "texte"
                    try:
                        with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
                            file_text = f.read(8000)
                        prompt = f"Analyse ce fichier texte ({filename}):\n\n{file_text}"
                        if message:
                            prompt += f"\n\nQuestion spécifique: {message}"
                        response_text, _ = get_gpt_response(prompt, max_tokens=1500, use_search=False, conversation_history=conversation_history)
                        question = f"Fichier: {filename}" + (f" + Question: {message}" if message else "")
                    except Exception as e:
                        response_text = f"Impossible de lire le fichier: {str(e)}"

                else:
                    response_text = f"Format de fichier non supporté pour l'analyse: {filename}"
                        
            finally:
                # Nettoyer le fichier uploadé
                if os.path.exists(filepath):
                    os.unlink(filepath)
                    
        elif message:
            # Déterminer le prompt image : demande directe ou affirmation suite à une proposition
            image_prompt = None
            file_subject = None
            if should_generate_image(message):
                image_prompt = message
            elif is_image_followup(message, conversation_history):
                # Chercher la vraie demande d'image dans l'historique (pas juste le dernier msg)
                image_prompt = extract_image_prompt_from_history(conversation_history) or message

            if image_prompt:
                message_type = "image_generated"
                image_url, img_error = generate_image(image_prompt)
                if image_url:
                    response_text = f"__IMAGE_GENERATED__:{image_url}"
                else:
                    response_text = img_error or "La génération d'image a échoué. Réessaie avec une description plus précise."
            else:
                # Détection PDF / Word
                file_type, file_subject = should_generate_file(message)
                if file_type:
                    message_type = f"file_{file_type}"
                    fmt_label = {'pdf': 'PDF', 'word': 'Word', 'pptx': 'PowerPoint'}.get(file_type, 'document')
                    # GPT génère le titre propre ET le contenu en une seule requête
                    gpt_prompt = (
                        f"Génère un contenu complet, bien structuré et professionnel pour un document "
                        f"{fmt_label} sur le sujet suivant : {file_subject}.\n\n"
                        f"Format de réponse OBLIGATOIRE :\n"
                        f"TITRE: [titre court et professionnel, 6 mots maximum]\n\n"
                        f"[contenu complet avec titres ## et listes]\n\n"
                        f"Utilise des titres markdown (##), des listes, et un contenu riche. "
                        f"Ne dis pas que tu génères un fichier — fournis directement le contenu."
                    )
                    raw_response, _ = get_gpt_response(gpt_prompt, max_tokens=3000, use_search=False)
                    # Extraire le titre de la première ligne
                    import re as _re
                    lines = raw_response.split('\n')
                    if lines and lines[0].strip().startswith('TITRE:'):
                        doc_title = lines[0].replace('TITRE:', '').strip()
                        doc_title = _re.sub(r'[*#"\'\[\]]', '', doc_title).strip()[:80]
                        doc_content = '\n'.join(lines[1:]).lstrip('\n')
                    else:
                        doc_title = file_subject[:60]
                        doc_title = _re.sub(r'[*\[\]]', '', doc_title).strip()
                        doc_content = raw_response
                    if not doc_title or len(doc_title) < 2:
                        doc_title = file_subject[:60]
                    response_text = (
                        f"__FILE_GENERATED__:{file_type}:{doc_title}\n{doc_content}"
                    )
                else:
                    response_text, sources = get_gpt_response(message, max_tokens=1500, conversation_history=conversation_history)
        else:
            response_text = "Veuillez saisir un message ou uploader un fichier."
        
        # Sauvegarder la conversation
        user_id = session['user_id']
        save_conversation(session_id, user_id, "web-interface", message_type, question, response_text)
        
        # Mettre à jour le titre de la session si c'est le premier message
        if len(conversation_history) == 0 and question:
            title = question[:50] + "..." if len(question) > 50 else question
            update_session_title(session_id, title)
        
        # Extraire l'URL si image générée
        image_url_out = None
        file_token_out = None

        if response_text.startswith('__IMAGE_GENERATED__:'):
            image_url_out = response_text.split(':', 1)[1]
            response_text = f"Voici l'image générée pour : *{message}*"

        elif response_text.startswith('__FILE_GENERATED__:'):
            # Format : __FILE_GENERATED__:pdf:Titre\nContenu...
            header, _, doc_content = response_text.partition('\n')
            parts = header.split(':', 2)
            file_type_out = parts[1] if len(parts) > 1 else 'pdf'
            doc_title = parts[2] if len(parts) > 2 else 'document'
            # Stocker le contenu en session Flask pour le téléchargement
            session['pending_file'] = {
                'type': file_type_out,
                'title': doc_title,
                'content': doc_content,
                'subject': file_subject or doc_title
            }
            file_token_out = file_type_out
            ext_map = {'pdf': 'PDF', 'pptx': 'PowerPoint', 'word': 'Word'}
            ext_label = ext_map.get(file_type_out, 'Word')
            ext_file = {'pdf': 'pdf', 'pptx': 'pptx'}.get(file_type_out, 'docx')
            response_text = (
                f"Votre document **{doc_title}** est prêt.\n\n"
                f"Cliquez sur le bouton ci-dessous pour le télécharger en {ext_label} (`.{ext_file}`)."
            )

        return jsonify({
            'success': True,
            'response': response_text,
            'type': message_type,
            'image_url': image_url_out,
            'file_ready': file_token_out,
            'sources': sources
        })
        
    except Exception as e:
        return jsonify({
            'success': False,
            'error': str(e)
        }), 500

@web_bp.route('/new-chat')
@login_required
def new_chat():
    """Crée une nouvelle session de chat"""
    user_id = session['user_id']
    session_id = create_chat_session(user_id)
    return redirect(f'/chat/{session_id}')

@web_bp.route('/download-file')
@login_required
def download_file():
    """Télécharge le dernier fichier PDF ou Word généré"""
    pending = session.get('pending_file')
    if not pending:
        return "Aucun fichier disponible.", 404

    file_type = pending['type']
    doc_title = pending['title']
    doc_content = pending['content']

    # Nettoyer le titre pour le nom de fichier
    import re
    safe_title = re.sub(r'[^\w\s-]', '', doc_title).strip().replace(' ', '_')[:60] or 'document'

    try:
        if file_type == 'pdf':
            file_bytes = generate_pdf(doc_content, doc_title)
            mimetype = 'application/pdf'
            filename = f"{safe_title}.pdf"
        elif file_type == 'pptx':
            from services.gpt_service import generate_image as _gen_img
            from services.file_generator import parse_sections as _parse_sec
            from concurrent.futures import ThreadPoolExecutor, as_completed
            import requests as _req

            subject = pending.get('subject', doc_title)

            # Toutes les sections de contenu
            all_sections = _parse_sec(doc_content)
            content_sections = [(i, t, b) for i, (t, b) in enumerate(all_sections)
                                if t and b.strip()]

            def _fetch(prompt, size="1792x1024"):
                try:
                    url, _ = _gen_img(prompt, size=size)
                    if url:
                        r = _req.get(url, timeout=30)
                        if r.status_code == 200:
                            return r.content
                except Exception:
                    pass
                return None

            cover_image = None
            slide_images = [None] * len(content_sections)

            # Génération parallèle : couverture + une image par slide
            with ThreadPoolExecutor(max_workers=5) as executor:
                futures = {}

                # Image de couverture
                cover_prompt = (
                    f"Professional corporate presentation cover photo about: {subject}. "
                    f"High-quality, modern business aesthetic, dark navy background with "
                    f"subtle purple and cyan light accents, no text, cinematic style."
                )
                futures['cover'] = executor.submit(_fetch, cover_prompt, "1024x1024")

                # Image par slide de contenu (max 8)
                for rank, (idx, sec_title, _) in enumerate(content_sections[:8]):
                    slide_prompt = (
                        f"Professional high-quality photo illustrating '{sec_title}' "
                        f"in the context of {subject}. "
                        f"Clean modern corporate style, sharp and bright, no text overlay, "
                        f"suitable for a business presentation slide."
                    )
                    futures[('slide', rank)] = executor.submit(_fetch, slide_prompt, "1792x1024")

                for key, future in futures.items():
                    try:
                        result = future.result(timeout=60)
                        if key == 'cover':
                            cover_image = result
                        else:
                            _, rank = key
                            slide_images[rank] = result
                    except Exception:
                        pass

            file_bytes = generate_pptx(
                doc_content, doc_title,
                cover_image=cover_image,
                slide_images=slide_images
            )
            mimetype = 'application/vnd.openxmlformats-officedocument.presentationml.presentation'
            filename = f"{safe_title}.pptx"
        else:
            file_bytes = generate_word(doc_content, doc_title)
            mimetype = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            filename = f"{safe_title}.docx"

        session.pop('pending_file', None)
        return send_file(
            io.BytesIO(file_bytes),
            mimetype=mimetype,
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        return f"Erreur lors de la génération du fichier : {str(e)}", 500

@web_bp.route('/voice-chat/<session_id>', methods=['POST'])
@login_required
def voice_chat(session_id):
    """
    Reçoit un blob audio (webm/mp3), transcrit avec Whisper,
    génère une réponse vocale GPT-4o, synthétise avec OpenAI TTS nova,
    retourne JSON { transcript, response_text, audio_b64 }.
    """
    import tempfile, base64

    audio_file = request.files.get('audio')
    if not audio_file:
        return jsonify({'success': False, 'error': 'Aucun fichier audio reçu.'}), 400

    # Détecter l'extension — Whisper accepte webm, mp4, mp3, wav, ogg
    filename = audio_file.filename or 'voice.webm'
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else 'webm'
    if ext not in ('webm', 'mp4', 'mp3', 'wav', 'ogg', 'm4a'):
        ext = 'webm'

    # Sauvegarder dans un fichier temporaire
    tmp_path = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=f'.{ext}') as tmp:
            audio_file.save(tmp)
            tmp_path = tmp.name

        # 1. Transcription Whisper
        try:
            lang = request.form.get('language') or None
            # Normaliser : "fr-FR" → "fr", "en-US" → "en"
            if lang and '-' in lang:
                lang = lang.split('-')[0]
            transcript = transcribe_audio(tmp_path, language=lang)
        except Exception as e:
            return jsonify({'success': False, 'error': f'Whisper: {str(e)}'}), 500

        if not transcript or transcript.startswith('Erreur'):
            return jsonify({'success': False, 'error': transcript or "Transcription échouée."}), 400

        # 2. Réponse GPT vocal
        try:
            conversation_history = get_conversation_history(session_id, 6)
            response_text = get_voice_response(transcript, conversation_history)
        except Exception as e:
            return jsonify({'success': False, 'error': f'GPT: {str(e)}'}), 500

        # 3. Synthèse TTS
        try:
            audio_bytes = tts_to_bytes(response_text, voice="nova")
            audio_b64 = base64.b64encode(audio_bytes).decode('utf-8') if audio_bytes else None
        except Exception as e:
            audio_b64 = None  # la voix est optionnelle — on continue sans

        # 4. Sauvegarder
        try:
            user_id = session['user_id']
            save_conversation(session_id, user_id, "web-voice", "voice", transcript, response_text)
            if len(conversation_history) == 0:
                update_session_title(session_id, transcript[:50])
        except Exception:
            pass  # non bloquant

        return jsonify({
            'success': True,
            'transcript': transcript,
            'response_text': response_text,
            'audio_b64': audio_b64
        })

    except Exception as e:
        return jsonify({'success': False, 'error': f'Erreur serveur: {str(e)}'}), 500
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except Exception:
                pass


@web_bp.route('/delete-session/<session_id>', methods=['POST'])
@login_required
def delete_session(session_id):
    """Supprime une session de chat"""
    user_id = session['user_id']
    try:
        delete_chat_session(user_id, session_id)
        return jsonify({'success': True})
    except Exception as e:
        return jsonify({'success': False, 'error': str(e)}), 500

@web_bp.route('/history')
@login_required
def history():
    """Affiche l'historique des conversations"""
    conversations = get_conversations(100)
    return render_template('history.html', conversations=conversations)